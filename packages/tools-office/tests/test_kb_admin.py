"""知识库后台文档管理（kb_admin）单元测试（原语直调，不起 HTTP 服务）。

覆盖：全格式入盘解析（md/txt/csv/docx/xlsx）、重传覆盖、来源清单三态
      （indexed / unindexed / missing）、向量化三口径（milvus / on-demand / unconfigured）
      与按源删除连向量一并清理、入参守卫（穿越/非白名单后缀/空体/超限/非法可见范围）。
固件隔离：monkeypatch settings.KB_DIR 到 tmp_path；默认钉空 EMBEDDING_*/MILVUS_URI/REDIS_URL
          零网络；向量库用内存 FakeVectorStore（实现 VectorStore Protocol）。
对齐：AGENTS.md §3（降级绝不 500 / 数据不出域）、§5（验证命令）；
      智能办公Agent 产品需求文档.md §2.13（知识库后台）、§2.6（知识库增强检索）。
"""

from __future__ import annotations

from pathlib import Path

import pytest

from office_agent_core import kv
from office_agent_core.contracts import ToolContext
from office_agent_core.errors import BusinessError
from office_agent_core.settings import settings
from office_agent_tools_office import kb, kb_admin, retrieval, vector_store

CTX = ToolContext(tenant="t1", username="alice", roles=["office:read"])


class FakeVectorStore:
    """内存版 VectorStore（与 test_vector_store 同口径，零网络、确定性可断言）。"""

    def __init__(self) -> None:
        self.rows: dict[str, vector_store.VectorRecord] = {}

    def available(self) -> bool:
        return True

    async def known_ids(self, ids: list[str]) -> set[str]:
        return {value for value in ids if value in self.rows}

    async def upsert(self, records: list[vector_store.VectorRecord]) -> int:
        for record in records:
            self.rows[record.id] = record
        return len(records)

    async def search(self, query_vector, top_k, min_score, origins=None):  # type: ignore[no-untyped-def]
        raise NotImplementedError("本用例集只验证灌库与删除，不验证检索排序")

    async def delete_by_source(self, origin: str, source: str) -> int:
        victims = [
            key
            for key, record in self.rows.items()
            if record.origin == origin and record.source == source
        ]
        for key in victims:
            self.rows.pop(key, None)
        return len(victims)


@pytest.fixture()
def kb_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """测试隔离的知识库目录（kb_admin 经 settings.KB_DIR 实时读取）。"""
    root = tmp_path / "kb"
    root.mkdir()
    monkeypatch.setattr(settings, "KB_DIR", str(root))
    return root


@pytest.fixture(autouse=True)
def _offline(monkeypatch: pytest.MonkeyPatch) -> None:
    """默认钉住「未配向量检索/向量库」：零网络走字符通道（本机 .env 一登记即会打网络）。"""
    monkeypatch.setattr(settings, "EMBEDDING_BASE_URL", "")
    monkeypatch.setattr(settings, "EMBEDDING_MODEL", "")
    monkeypatch.setattr(settings, "MILVUS_URI", "")
    monkeypatch.setattr(settings, "REDIS_URL", "")
    vector_store.reset_store()
    kv.reset()
    yield
    vector_store.reset_store()
    kv.reset()


async def _fake_embed(texts: list[str]) -> list[list[float]]:
    """确定性假向量（长度归一，同文本同向量），替代真实 /embeddings。"""
    return [[len(text) / 100.0, 1.0] for text in texts]


def _enable_embedding(monkeypatch: pytest.MonkeyPatch, store: FakeVectorStore | None) -> None:
    """配好向量通道：EMBEDDING_* 有值 + embed_texts 打桩 + 向量库替身（None=不配库）。"""
    monkeypatch.setattr(settings, "EMBEDDING_BASE_URL", "http://127.0.0.1:11434/v1")
    monkeypatch.setattr(settings, "EMBEDDING_MODEL", "fake-embedding")
    monkeypatch.setattr(retrieval, "embed_texts", _fake_embed)
    monkeypatch.setattr(vector_store, "get_store", lambda: store)


# ---------------- ① 全格式入盘 + 解析 + 清单 ----------------

MD_DOC = "报销制度\n费用发生后 30 天内提交，单笔超 1000 元需部门负责人审批。"


async def test_save_and_ingest_parses_and_marks_manifest(kb_dir: Path) -> None:
    """上传即入盘、解析、切块，并把元数据写进来源清单（可见范围与上传人如实记账）。"""
    result = await kb_admin.save_and_ingest(
        "制度.md", MD_DOC.encode("utf-8"), uploaded_by="admin", visibility="hr"
    )
    assert result["filename"] == "制度.md"
    assert result["format"] == "md"
    assert result["title"] == "报销制度"
    assert result["chunks"] == 1
    assert result["degraded"] is False
    assert result["reuploaded"] is False
    assert result["vector_mode"] == "unconfigured"  # 未配 EMBEDDING_*：如实说没向量化
    assert (kb_dir / "制度.md").is_file()

    listing = await kb_admin.list_sources()
    assert listing["total"] == 1
    row = listing["items"][0]
    assert row["state"] == "indexed"
    assert row["state_label"] == "已入库"
    assert row["visibility"] == "hr"
    assert row["uploaded_by"] == "admin"
    assert row["chunks"] == 1
    assert listing["summary"] == {
        "files": 1,
        "indexed": 1,
        "unindexed": 0,
        "missing": 0,
        "chunks": 1,
        "vectorized_files": 0,
        "vectorized_chunks": 0,
        "total_bytes": len(MD_DOC.encode("utf-8")),
    }


async def test_ingests_text_and_csv_formats(kb_dir: Path) -> None:
    """TXT / Markdown / CSV 走同一提取口径，全部可解析入库。"""
    await kb_admin.save_and_ingest(
        "a.txt", "考勤制度\n工作日 9:00-18:00。".encode(), uploaded_by="admin", visibility="public"
    )
    await kb_admin.save_and_ingest(
        "b.md",
        "# 差旅标准\n高铁二等座为默认标准。".encode(),
        uploaded_by="admin",
        visibility="public",
    )
    await kb_admin.save_and_ingest(
        "c.csv", "部门,金额\n销售,1200\n".encode(), uploaded_by="admin", visibility="public"
    )
    listing = await kb_admin.list_sources()
    by_name = {item["filename"]: item for item in listing["items"]}
    assert set(by_name) == {"a.txt", "b.md", "c.csv"}
    assert all(item["state"] == "indexed" for item in by_name.values())
    assert by_name["b.md"]["title"] == "差旅标准"  # markdown 修饰符被剥掉
    assert all(item["chunks"] >= 1 for item in by_name.values())


async def test_ingests_excel_and_word(kb_dir: Path) -> None:
    """Excel / Word 二进制格式同样走 file_read 提取口径（依赖缺失则如实降级不伪装）。"""
    from docx import Document
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.append(["项目", "负责人"])
    sheet.append(["上线培训", "张三"])
    book.save(kb_dir / "台账.xlsx")

    doc = Document()
    doc.add_paragraph("会议纪要")
    doc.add_paragraph("决议：本周五前完成接口联调。")
    doc.save(kb_dir / "纪要.docx")

    for name in ("台账.xlsx", "纪要.docx"):
        result = await kb_admin.ingest_source(name, uploaded_by="admin", visibility="public")
        assert result["format"] in ("xlsx", "docx")
        assert result["degraded"] is False, result["degraded_reason"]
        assert result["chunks"] >= 1
    listing = await kb_admin.list_sources()
    assert {item["filename"] for item in listing["items"]} == {"台账.xlsx", "纪要.docx"}


def test_supported_formats_cover_prd_list() -> None:
    """PRD §2.13 要求的 PDF / Word / Excel / TXT / Markdown 全在白名单内（单一出处）。"""
    assert set(kb_admin.KB_SUFFIXES) == {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md"}


# ---------------- ② 入参守卫 ----------------


async def test_save_rejects_traversal_suffix_empty_and_bad_visibility(kb_dir: Path) -> None:
    """穿越 / 非白名单后缀 / 空体 / 非法可见范围一律 1001（不落盘）。"""
    for filename, data, visibility in (
        ("../evil.md", b"x", "public"),
        ("sub/evil.md", b"x", "public"),
        ("evil.exe", b"x", "public"),
        ("a.md", b"", "public"),
        ("a.md", b"x", "has space"),
    ):
        with pytest.raises(BusinessError) as exc:
            await kb_admin.save_and_ingest(
                filename, data, uploaded_by="admin", visibility=visibility
            )
        assert exc.value.code == 1001, (filename, visibility)
    assert (await kb_admin.list_sources())["total"] == 0  # 一律未落盘


async def test_save_rejects_oversize_without_writing(
    kb_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """超上限即 1001 且不落盘（上限与解析上限同一出处，杜绝「已上传未入库」）。"""
    monkeypatch.setattr(kb_admin, "MAX_UPLOAD_BYTES", 8)
    with pytest.raises(BusinessError) as exc:
        await kb_admin.save_and_ingest(
            "a.md", b"123456789", uploaded_by="admin", visibility="public"
        )
    assert exc.value.code == 1001
    assert (await kb_admin.list_sources())["total"] == 0


async def test_delete_unknown_source_is_404(kb_dir: Path) -> None:
    with pytest.raises(BusinessError) as exc:
        await kb_admin.delete_source("nope.md")
    assert exc.value.code == 1004


# ---------------- ③ 来源清单三态（不掩盖异常） ----------------


async def test_manual_file_on_disk_is_unindexed_then_indexable(kb_dir: Path) -> None:
    """手工放进目录的文件如实标「已入盘未解析」，补解析后转「已入库」。"""
    (kb_dir / "手工.md").write_text("手工放入的资料\n未被后台上传过。", encoding="utf-8")
    listing = await kb_admin.list_sources()
    assert listing["items"][0]["state"] == "unindexed"
    assert listing["items"][0]["chunks"] == 0
    assert listing["summary"]["unindexed"] == 1

    result = await kb_admin.ingest_source("手工.md", uploaded_by="admin", visibility="public")
    assert result["chunks"] == 1
    listing = await kb_admin.list_sources()
    assert listing["items"][0]["state"] == "indexed"
    assert listing["summary"]["indexed"] == 1


async def test_missing_file_still_listed_and_deletable(kb_dir: Path) -> None:
    """清单在但文件被外部删掉 → 如实标「文件缺失」，且仍可按源清理残留条目。"""
    await kb_admin.save_and_ingest(
        "gone.md", MD_DOC.encode(), uploaded_by="admin", visibility="public"
    )
    (kb_dir / "gone.md").unlink()
    listing = await kb_admin.list_sources()
    assert listing["items"][0]["state"] == "missing"
    assert listing["summary"] == {
        "files": 0,
        "indexed": 0,
        "unindexed": 0,
        "missing": 1,
        "chunks": 0,
        "vectorized_files": 0,
        "vectorized_chunks": 0,
        "total_bytes": 0,
    }
    result = await kb_admin.delete_source("gone.md")
    assert result["deleted"] is True
    assert (await kb_admin.list_sources())["total"] == 0


async def test_unparsable_file_degrades_but_stays_traceable(
    kb_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """解析失败（引擎缺失/文件损坏）如实降级：入盘可查、切片 0、原因写清，绝不 500。"""
    (kb_dir / "破损.pdf").write_bytes(b"not a real pdf")
    result = await kb_admin.ingest_source("破损.pdf", uploaded_by="admin", visibility="public")
    assert result["degraded"] is True
    assert result["chunks"] == 0
    assert result["degraded_reason"]
    row = (await kb_admin.list_sources())["items"][0]
    assert row["state"] == "indexed"  # 已入盘（磁盘有文件）
    assert row["degraded"] is True
    assert row["degraded_reason"] == result["degraded_reason"]


# ---------------- ④ 向量化三口径 + 按源删除 ----------------


async def test_vectorize_milvus_ingest_and_delete_clears_vectors(
    kb_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """配向量库：上传即落库（vectorized=True/mode=milvus），删除时该源向量一并清理。"""
    store = FakeVectorStore()
    _enable_embedding(monkeypatch, store)

    result = await kb_admin.save_and_ingest(
        "制度.md", MD_DOC.encode(), uploaded_by="admin", visibility="public"
    )
    assert result["vector_mode"] == "milvus"
    assert result["vectorized"] is True
    assert result["vector_reason"].startswith("已落库")
    assert len(store.rows) == result["chunks"] == 1
    assert next(iter(store.rows.values())).origin == kb_admin.ORIGIN
    assert next(iter(store.rows.values())).source == "local-kb:制度.md"

    stats = await kb_admin.kb_stats()
    assert stats["summary"]["vectorized_files"] == 1
    assert stats["summary"]["vectorized_chunks"] == 1
    assert stats["store"]["mode"] == "milvus"
    assert stats["embedding_model"] == "fake-embedding"

    deleted = await kb_admin.delete_source("制度.md")
    assert deleted["removed_vectors"] == 1
    assert store.rows == {}
    assert not (kb_dir / "制度.md").exists()
    assert (await kb_admin.list_sources())["total"] == 0


async def test_reupload_clears_stale_vectors(kb_dir: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """重传同名文件：先清该源旧向量再重新入库，旧文本块不留库里积垃圾。"""
    store = FakeVectorStore()
    _enable_embedding(monkeypatch, store)
    await kb_admin.save_and_ingest(
        "制度.md", MD_DOC.encode(), uploaded_by="admin", visibility="public"
    )
    old_ids = set(store.rows)

    again = await kb_admin.save_and_ingest(
        "制度.md",
        "报销制度（已修订）\n期限由 30 天改为 15 天。".encode(),
        uploaded_by="admin",
        visibility="public",
    )
    assert again["reuploaded"] is True
    assert set(store.rows) & old_ids == set()  # 旧块已清
    assert len(store.rows) == 1
    assert "15 天" in next(iter(store.rows.values())).text


async def test_vectorize_on_demand_without_milvus(
    kb_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """配了 EMBEDDING_* 但没配向量库：如实标「检索时按需算、不落库」，绝不谎称已向量化。"""
    _enable_embedding(monkeypatch, None)
    result = await kb_admin.save_and_ingest(
        "制度.md", MD_DOC.encode(), uploaded_by="admin", visibility="public"
    )
    assert result["vectorized"] is False
    assert result["vector_mode"] == "on_demand_embedding"
    assert "按需计算" in result["vector_reason"]
    stats = await kb_admin.kb_stats()
    assert stats["store"]["mode"] == "on_demand_embedding"
    assert stats["embedding_model"] == "fake-embedding"  # 语义检索可用，只是不落库


async def test_vector_store_failure_degrades_without_raising(
    kb_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """向量库对端故障：降级为按需向量化并如实标注，文件仍正常入库（降级不 500 红线）。"""

    class _BrokenStore(FakeVectorStore):
        async def known_ids(self, ids: list[str]) -> set[str]:
            raise ConnectionError("milvus down")

    _enable_embedding(monkeypatch, _BrokenStore())
    result = await kb_admin.save_and_ingest(
        "制度.md", MD_DOC.encode(), uploaded_by="admin", visibility="public"
    )
    assert result["degraded"] is False  # 入库本身成功
    assert result["vectorized"] is False
    assert result["vector_mode"] == "on_demand_embedding"
    assert "Milvus" in result["vector_reason"] or "milvus" in result["vector_reason"].lower()


# ---------------- ⑤ 统计与检索口径贯通 ----------------


async def test_stats_groups_by_format(kb_dir: Path) -> None:
    """统计按格式分组（文件数/切片数/字节数三者同源自真实清点）。"""
    await kb_admin.save_and_ingest(
        "a.md", MD_DOC.encode(), uploaded_by="admin", visibility="public"
    )
    await kb_admin.save_and_ingest(
        "b.md", MD_DOC.encode(), uploaded_by="admin", visibility="public"
    )
    await kb_admin.save_and_ingest(
        "c.txt", "考勤制度\n工作日 9:00-18:00。".encode(), uploaded_by="admin", visibility="public"
    )
    stats = await kb_admin.kb_stats()
    buckets = {item["format"]: item for item in stats["by_format"]}
    assert buckets["md"] == {
        "format": "md",
        "files": 2,
        "chunks": 2,
        "bytes": 2 * len(MD_DOC.encode()),
    }
    assert buckets["txt"]["files"] == 1
    assert stats["summary"]["files"] == 3
    assert stats["supported_formats"] == [suffix.lstrip(".") for suffix in kb_admin.KB_SUFFIXES]


async def test_uploaded_file_is_immediately_retrievable(kb_dir: Path) -> None:
    """上传即可被 kb.ask 检索到（同一提取/切块/来源口径，无需额外注册步骤）。"""
    await kb_admin.save_and_ingest(
        "远程办公制度.md",
        "远程办公制度\n每周最多 2 天远程办公，需提前一天在系统报备。".encode(),
        uploaded_by="admin",
        visibility="public",
    )
    data = await kb._kb_ask(CTX, {"query": "远程办公每周最多几天", "top_k": 1})
    assert data["count"] == 1
    assert data["results"][0]["source"] == "local-kb:远程办公制度.md"
    assert data["results"][0]["visibility"] == "public"
    assert "2 天" in data["results"][0]["snippet"]


async def test_manifest_visibility_is_enforced_for_binary_formats(kb_dir: Path) -> None:
    """受限来源靠来源清单生效权限过滤：docx 正文写不了 visibility 指令，只能由清单定权。"""
    from docx import Document

    doc = Document()
    doc.add_paragraph("股权激励归属办法")
    doc.add_paragraph("四年分批归属，每年归属 25%。")
    doc.save(kb_dir / "股权.docx")
    await kb_admin.ingest_source("股权.docx", uploaded_by="admin", visibility="hr")
    assert kb_admin.manifest_visibility() == {"股权.docx": "hr"}

    outsider = ToolContext(tenant="t1", username="bob", roles=["office:read"])
    blocked = await kb._kb_ask(outsider, {"query": "股权激励怎么归属"})
    assert blocked["permission_filtered"] >= 1
    assert all(hit["source"] != "local-kb:股权.docx" for hit in blocked["results"])

    insider = ToolContext(tenant="t1", username="hr01", roles=["office:read", "hr"])
    allowed = await kb._kb_ask(insider, {"query": "股权激励怎么归属", "top_k": 1})
    assert allowed["results"][0]["source"] == "local-kb:股权.docx"
    assert allowed["results"][0]["visibility"] == "hr"
