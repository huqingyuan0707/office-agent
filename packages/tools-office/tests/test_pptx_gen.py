"""PPT 生成工具单元测试（handler 直调，不起 HTTP 服务）。

覆盖：大纲直出写盘（python-pptx 可读回）、参数校验（空大纲/越界/穿越拒绝）、
      ToolSpec 写口径（恒送审 + 幂等键必带）与注册聚合。
对齐：AGENTS.md §5（验证命令）；智能办公Agent 产品需求文档.md §2.12（V1.2 PPT 生成）。
"""

from __future__ import annotations

from pathlib import Path

import pytest

from office_agent_core.contracts import ToolContext
from office_agent_core.errors import BusinessError
from office_agent_core.settings import settings
from office_agent_tools_office import pptx_gen

CTX = ToolContext(tenant="t1", username="alice", roles=["office:write"])

_SLIDES = [
    {"title": "本周进展", "bullets": ["日报助手上线", "审批闭环跑通"]},
    {"title": "下周计划", "bullets": ["V1.2 批次 B 前端两页"]},
]


def _read_deck(path: Path) -> list[str]:
    """读回 pptx 文本（校验内容确为入参原值直出）。"""
    from pptx import Presentation

    deck = Presentation(str(path))
    texts: list[str] = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    return texts


# ---------------- handler 行为 ----------------


async def test_generate_writes_deck(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "DOCS_DIR", str(tmp_path))
    data = await pptx_gen._pptx_generate(
        CTX, {"title": "周报复盘", "filename": "weekly.pptx", "slides": _SLIDES}
    )
    assert data["published"] is True
    assert data["slide_count"] == 2
    assert data["file"] == "weekly.pptx"
    texts = "\n".join(_read_deck(tmp_path / "weekly.pptx"))
    assert "周报复盘" in texts
    assert "日报助手上线" in texts
    assert "V1.2 批次 B 前端两页" in texts


async def test_generate_rejects_empty_slides(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "DOCS_DIR", str(tmp_path))
    try:
        await pptx_gen._pptx_generate(CTX, {"title": "周报", "filename": "a.pptx", "slides": []})
    except BusinessError as exc:
        assert exc.code == 1001
    else:
        raise AssertionError("空大纲应被 1001 拒绝")


async def test_generate_rejects_blank_bullet(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "DOCS_DIR", str(tmp_path))
    try:
        await pptx_gen._pptx_generate(
            CTX,
            {"title": "周报", "filename": "a.pptx", "slides": [{"title": "页", "bullets": ["  "]}]},
        )
    except BusinessError as exc:
        assert exc.code == 1001
    else:
        raise AssertionError("空白要点应被 1001 拒绝")


async def test_generate_rejects_traversal(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "DOCS_DIR", str(tmp_path))
    try:
        await pptx_gen._pptx_generate(
            CTX, {"title": "周报", "filename": "../evil.pptx", "slides": _SLIDES}
        )
    except BusinessError as exc:
        assert exc.code == 1001
    else:
        raise AssertionError("路径穿越应被拒绝")


# ---------------- 注册口径 ----------------


def test_spec_is_write_with_approval_and_idem() -> None:
    (spec,) = pptx_gen.specs()
    assert spec.name == "office.pptx.generate"
    assert spec.scope == "office:write"
    assert spec.requires_approval is True
    assert spec.idempotent is True
    assert spec.approval_action == "office.pptx.generate"
    assert "idem_key" in spec.params["required"]
