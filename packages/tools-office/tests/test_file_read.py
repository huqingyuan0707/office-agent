"""文件解析与批量重命名单测（handler 直调，不起 HTTP 服务）。

覆盖：docx 段落 + 表格文本 + 图片清单、xlsx 行列与文本、txt 直读、PDF 真实逐页抽取与
      引擎缺失降级、缺文件 404、穿越拒绝、批量重命名成功/禁覆盖/源缺失/注册口径。
固件隔离：monkeypatch settings.DOCS_DIR 到 tmp_path，仓库 data 目录零污染。
对齐：AGENTS.md §5（验证命令）；智能办公Agent 产品需求文档.md §2.1（文件解析/批量处理）。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from PIL import Image

from office_agent_core.contracts import ToolContext
from office_agent_core.errors import BusinessError
from office_agent_core.settings import settings
from office_agent_tools_office import file_read

CTX = ToolContext(tenant="t1", username="alice", roles=["office:read", "office:write"])


@pytest.fixture()
def docs_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """测试隔离的文档工作目录（handler 经 settings.DOCS_DIR 实时读取）。"""
    monkeypatch.setattr(settings, "DOCS_DIR", str(tmp_path))
    return tmp_path


def _make_docx(path: Path, with_image: Path | None = None) -> None:
    """造 docx：两段 + 1 表 + 可选 1 图。"""
    document = Document()
    document.add_paragraph("第一段：项目背景介绍")
    document.add_paragraph("第二段：本期工作进展")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "张三"
    if with_image is not None:
        Image.new("RGB", (8, 8), color=(200, 200, 200)).save(with_image, format="PNG")
        document.add_picture(str(with_image), width=Inches(0.5))
    document.save(str(path))


def _make_pdf(path: Path, text: str, pages: int = 1) -> None:
    """手写最小 PDF（含真实计算的 xref 偏移）：N 页、每页一行文本。

    结构化生成而非引第三方库（reportlab 非本项目依赖）；偏移量必须准确，否则 pypdf 解析失败。
    """
    page_ids = [4 + index for index in range(pages)]
    content_ids = [4 + pages + index for index in range(pages)]
    kids = b" ".join(f"{page_id} 0 R".encode() for page_id in page_ids)
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(pages).encode() + b" >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for page_index in range(pages):
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents "
            + str(content_ids[page_index]).encode()
            + b" 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        )
    for page_index in range(pages):
        stream = f"BT /F1 12 Tf 20 100 Td (page {page_index + 1}: {text}) Tj ET".encode("ascii")
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    ).encode()
    path.write_bytes(bytes(out))


# ---------------- 文件解析 ----------------


async def test_read_docx_paragraphs_table_images(docs_dir: Path) -> None:
    """docx：段落 + 表格文本 + 图片清单。"""
    _make_docx(docs_dir / "demo.docx", with_image=docs_dir / "pic.png")
    data = await file_read._file_read(CTX, {"filename": "demo.docx"})
    assert data["format"] == "docx"
    assert "项目背景介绍" in data["text"]
    assert "姓名 | 张三" in data["text"]
    assert len(data["images"]) == 1 and data["images"][0].endswith(".png")
    assert data["degraded"] is False
    assert data["source"] == "local-docs:demo.docx"


async def test_read_xlsx_sheets_and_text(docs_dir: Path) -> None:
    """xlsx：每表行列数 + 前 N 行文本。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "费用"
    sheet.append(["项目", "金额"])
    sheet.append(["差旅", 1200])
    workbook.save(str(docs_dir / "cost.xlsx"))
    data = await file_read._file_read(CTX, {"filename": "cost.xlsx"})
    assert data["sheets"] == [{"name": "费用", "rows": 2, "cols": 2}]
    assert "差旅 | 1200" in data["text"]


async def test_read_txt_direct(docs_dir: Path) -> None:
    """txt 直读 + 溯源。"""
    (docs_dir / "note.txt").write_text(" hello \n世界", encoding="utf-8")
    data = await file_read._file_read(CTX, {"filename": "note.txt"})
    assert "世界" in data["text"]
    assert data["source"] == "local-docs:note.txt"


async def test_read_pdf_extracts_pages(docs_dir: Path) -> None:
    """PDF 走 pypdf 真实抽取：文本与页数如实返回（不引入 reportlab，手写最小 PDF）。"""
    _make_pdf(docs_dir / "plan.pdf", "Office Agent PDF Check", pages=2)
    data = await file_read._file_read(CTX, {"filename": "plan.pdf"})
    assert data["format"] == "pdf"
    assert data["degraded"] is False
    assert data["pages"] == 2
    assert "Office Agent PDF Check" in data["text"]
    assert "第 1 页" in data["text"]


async def test_read_pdf_engine_missing_degrades(
    docs_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """引擎缺失如实降级（不编造半页文字）——按「pypdf 未安装」模拟。"""
    monkeypatch.setattr(file_read, "_PDF_AVAILABLE", False)
    (docs_dir / "scan.pdf").write_bytes(b"%PDF-1.4 fake")
    data = await file_read._file_read(CTX, {"filename": "scan.pdf"})
    assert data["degraded"] is True
    assert "pypdf" in data["degraded_reason"]
    assert data["text"] == ""


async def test_read_pdf_broken_file_rejects(docs_dir: Path) -> None:
    """损坏 PDF → 1001 可操作提示（不 500、不假装读出内容）。"""
    (docs_dir / "broken.pdf").write_bytes(b"not a pdf at all")
    try:
        await file_read._file_read(CTX, {"filename": "broken.pdf"})
    except BusinessError as exc:
        assert exc.code == 1001
    else:
        raise AssertionError("损坏 PDF 应被 1001 拒绝")


async def test_read_missing_file_404(docs_dir: Path) -> None:
    """缺文件 404。"""
    _ = docs_dir
    try:
        await file_read._file_read(CTX, {"filename": "gone.docx"})
    except BusinessError as exc:
        assert exc.code == 1004
    else:
        raise AssertionError("缺文件应 404")


async def test_read_rejects_traversal(docs_dir: Path) -> None:
    """穿越拒绝。"""
    _ = docs_dir
    try:
        await file_read._file_read(CTX, {"filename": "../evil.docx"})
    except BusinessError as exc:
        assert exc.code == 1001
    else:
        raise AssertionError("穿越应被 1001 拒绝")


# ---------------- 批量重命名 ----------------


async def test_rename_batch_ok(docs_dir: Path) -> None:
    """批量重命名成功（源消失目标出现）。"""
    (docs_dir / "a.txt").write_text("a", encoding="utf-8")
    (docs_dir / "b.txt").write_text("b", encoding="utf-8")
    data = await file_read._docs_rename(
        CTX, {"pairs": [{"src": "a.txt", "dst": "a2.txt"}, {"src": "b.txt", "dst": "b2.txt"}]}
    )
    assert data["count"] == 2
    assert (docs_dir / "a2.txt").is_file() and not (docs_dir / "a.txt").exists()


async def test_rename_refuses_overwrite_and_missing(docs_dir: Path) -> None:
    """目标已存在禁覆盖；源缺失 404 诚实失败。"""
    (docs_dir / "x.txt").write_text("x", encoding="utf-8")
    (docs_dir / "y.txt").write_text("y", encoding="utf-8")
    try:
        await file_read._docs_rename(CTX, {"pairs": [{"src": "x.txt", "dst": "y.txt"}]})
    except BusinessError as exc:
        assert exc.code == 1001
    else:
        raise AssertionError("覆盖应被 1001 拒绝")
    try:
        await file_read._docs_rename(CTX, {"pairs": [{"src": "nope.txt", "dst": "n2.txt"}]})
    except BusinessError as exc:
        assert exc.code == 1004
    else:
        raise AssertionError("源缺失应 404")


def test_file_tools_register_scopes() -> None:
    """注册口径：读免审 / 写恒送审且非幂等（重放诚实失败）。"""
    specs = {spec.name: spec for spec in file_read.specs()}
    assert set(specs) == {"office.file.read", "office.docs.rename"}
    assert specs["office.file.read"].requires_approval is False
    assert specs["office.docs.rename"].requires_approval is True
    assert specs["office.docs.rename"].scope == "office:write"
    assert specs["office.docs.rename"].idempotent is False
