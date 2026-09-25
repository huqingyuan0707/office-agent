"""Word 直出工具单测（docx_render.py，handler 直调，不起 HTTP 服务）。

覆盖：markdown 子集渲染 docx 字节可解码可回读（标题/列表/表格/加粗往返）、入参校验
      （空标题/空正文/超长）、文件名非法字符清洗、注册口径（读免审不写盘）。
对齐：AGENTS.md §5（验证命令）；智能办公Agent 产品需求文档.md §2.1。
"""

from __future__ import annotations

import base64
import io

import pytest
from docx import Document

from office_agent_core.contracts import ToolContext
from office_agent_core.errors import BusinessError
from office_agent_tools_office import docx_render

CTX = ToolContext(tenant="t1", username="alice", roles=["office:read"])

_MD = (
    "# 工作日报\n"
    "生成时间：2026-09-25（模板直出）\n\n"
    "## 一、核心指标\n"
    "- 完成任务：**8** 项\n"
    "- 待办剩余：3 项\n\n"
    "## 二、明细\n"
    "| 事项 | 状态 |\n"
    "| --- | --- |\n"
    "| 日报 | 完成 |\n\n"
    "1. 第一项跟进\n"
    "2. 第二项跟进\n"
)


async def test_render_returns_decodable_docx_roundtrip() -> None:
    data = await docx_render._docx_render(CTX, {"title": "工作日报", "markdown": _MD})
    assert data["encoding"] == "base64"
    assert data["mime"] == docx_render.DOCX_MIME
    assert data["source"] == "office.docx.render" and data["fetched_at"]
    raw = base64.b64decode(data["content"])
    assert raw[:2] == b"PK" and data["size_bytes"] == len(raw)
    # 回读校验：与标题重复的首个一级标题被跳过，正文小标题/列表/表格/加粗都在
    text = "\n".join(p.text for p in Document(io.BytesIO(raw)).paragraphs)
    assert text.count("工作日报") == 1
    assert "一、核心指标" in text and "完成任务：8 项" in text
    table_texts = [c.text for row in Document(io.BytesIO(raw)).tables[0].rows for c in row.cells]
    assert "事项" in table_texts and "日报" in table_texts


async def test_render_rejects_blank_params() -> None:
    with pytest.raises(BusinessError):
        await docx_render._docx_render(CTX, {"title": " ", "markdown": _MD})
    with pytest.raises(BusinessError):
        await docx_render._docx_render(CTX, {"title": "T", "markdown": "\n  \n"})
    with pytest.raises(BusinessError):
        await docx_render._docx_render(CTX, {"title": "T", "markdown": "x" * 20001})


async def test_filename_sanitized_from_title() -> None:
    data = await docx_render._docx_render(
        CTX, {"title": 'a/b\\c:d*e?"f<g>h|i', "markdown": "正文内容"}
    )
    assert "/" not in data["filename"] and ":" not in data["filename"]
    assert data["filename"].endswith(".docx")


def test_spec_registered_read_no_approval() -> None:
    specs = docx_render.specs()
    assert len(specs) == 1
    spec = specs[0]
    assert spec.name == "office.docx.render"
    assert spec.scope == "office:read"
    assert spec.requires_approval is False
