"""文件解析与批量重命名工具（office.file.read / docs.rename，对齐 PRD §2.1 文件解析/批量处理）。

职责：
- office.file.read（office:read）：DOCS_DIR 内 docx/xlsx/csv/txt/md 内容提取——
  docx 取段落 + 表格文本 + 图片清单，xlsx 取每表行列数 + 前 N 行文本，纯文本直读；
  PDF 明确降级（pypdf 未安装：degraded + 原因，不编造半页文字）；缺文件 404；
- office.docs.rename（office:write + 恒送审）：DOCS_DIR 内批量重命名（pairs 上限 20 对，
  双向 basename 防穿越 + 后缀白名单 + 禁止覆盖已存在目标）；重放因源文件已迁走而
  诚实失败，故标 idempotent=False（不谎称幂等）。

链路：__init__.register_all() → registry.register(spec) → executor.call 执行 handler。
红线：纯本地实现，不触及 ORM / FastAPI；文件名一律锁 DOCS_DIR（paths.resolve_under_docs，
      与 tools_docs.py 同款口径）；解析引擎缺失只降级不阻断，绝不 500。
对齐：AGENTS.md §3（分层/数据不出域在本地盘的对应实现）；智能办公Agent 产品需求文档.md
      §2.1（文件解析：Word/PDF/Excel 提取——PDF 待引擎；批量处理：重命名/提图/批量摘要——
      提图并入 file.read 图片清单，批量摘要复用 office.text.summarize texts 多篇模式，
      批量转 PDF 需 soffice 外部二进制，明确后置，三者均不重复造）。
"""

from __future__ import annotations

import asyncio
import csv
import io
import logging
from pathlib import Path
from typing import Any

from office_agent_core.contracts import ToolContext, ToolSpec
from office_agent_core.errors import BusinessError, ErrorCode
from office_agent_core.registry import register
from office_agent_tools_office.paths import resolve_under_docs

logger = logging.getLogger(__name__)

SCOPE_READ = "office:read"
SCOPE_WRITE = "office:write"

_READ_SUFFIXES = (".docx", ".xlsx", ".csv", ".txt", ".md", ".pdf")
_RENAME_SUFFIXES = (".docx", ".xlsx", ".csv", ".txt", ".md", ".png", ".jpg", ".pptx")
_MAX_FILE_BYTES = 200_000
_MAX_TEXT_CHARS = 20000
_MAX_PAIRS = 20
_MAX_SHEET_ROWS = 200
_MAX_SHEET_COLS = 20

try:
    import docx as _docx_module

    _DOCX_AVAILABLE = True
except ImportError:
    _docx_module = None
    _DOCX_AVAILABLE = False

try:
    import openpyxl as _openpyxl_module

    _XLSX_AVAILABLE = True
except ImportError:
    _openpyxl_module = None
    _XLSX_AVAILABLE = False


def _truncate(text: str) -> tuple[str, bool]:
    """超长截断（截断标记明示，不静默丢尾）。"""
    if len(text) > _MAX_TEXT_CHARS:
        return text[:_MAX_TEXT_CHARS] + "…（已截断）", True
    return text, False


def _read_docx(path: Path) -> dict[str, Any]:
    """docx 提取：段落 + 表格文本 + 图片清单（media 内嵌件名）。"""
    assert _docx_module is not None
    document = _docx_module.Document(str(path))
    parts = [para.text.strip() for para in document.paragraphs if para.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    media: list[str] = []
    for rel in document.part.rels.values():
        target = rel.target_ref
        # 图片 rel 的 target_ref 是相对路径（media/xxx），非 word/media/ 全路径
        if isinstance(target, str) and "media/" in target:
            media.append(target.split("/")[-1])
    text, truncated = _truncate("\n".join(parts))
    return {"text": text, "truncated": truncated, "images": sorted(set(media))}


def _read_xlsx(path: Path) -> dict[str, Any]:
    """xlsx 提取：每表行列数 + 前 N 行文本（read_only 流式，超大表不爆内存）。

    openpyxl read_only 下 max_row/max_column 仍可用；逐行 iter_rows 取值。
    """
    assert _openpyxl_module is not None
    workbook = _openpyxl_module.load_workbook(str(path), read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    chunks: list[str] = []
    try:
        for name in workbook.sheetnames:
            sheet = workbook[name]
            rows = list(
                sheet.iter_rows(
                    min_row=1, max_row=_MAX_SHEET_ROWS, max_col=_MAX_SHEET_COLS, values_only=True
                )
            )
            sheets.append({"name": name, "rows": sheet.max_row or 0, "cols": sheet.max_column or 0})
            for row in rows:
                cells = [str(cell).strip() for cell in row if cell not in (None, "")]
                if cells:
                    chunks.append(" | ".join(cells))
    finally:
        workbook.close()
    text, truncated = _truncate("\n".join(chunks))
    return {"text": text, "truncated": truncated, "sheets": sheets}


def _read_textlike(path: Path) -> dict[str, Any]:
    """csv/txt/md 直读（csv 按行逗号转竖线拼可读文本，不做语义解析）。"""
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    if path.suffix.lower() == ".csv":
        try:
            rows = list(csv.reader(io.StringIO(raw)))[:_MAX_SHEET_ROWS]
            raw = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows if row)
        except csv.Error:
            pass
    text, truncated = _truncate(raw.strip())
    return {"text": text, "truncated": truncated}


def _extract(path: Path) -> dict[str, Any]:
    """按后缀分发提取（同步 IO，调用方包 asyncio.to_thread）。"""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return {
            "text": "",
            "truncated": False,
            "degraded": True,
            "degraded_reason": "PDF 解析引擎缺失（pypdf 未安装）：请先转存为 docx/txt 后再解析",
        }
    if suffix == ".docx":
        if not _DOCX_AVAILABLE:
            return {
                "text": "",
                "truncated": False,
                "degraded": True,
                "degraded_reason": "docx 解析依赖缺失（python-docx 未安装）",
            }
        return {"degraded": False, "degraded_reason": "", **_read_docx(path)}
    if suffix == ".xlsx":
        if not _XLSX_AVAILABLE:
            return {
                "text": "",
                "truncated": False,
                "degraded": True,
                "degraded_reason": "xlsx 解析依赖缺失（openpyxl 未安装）",
            }
        return {"degraded": False, "degraded_reason": "", **_read_xlsx(path)}
    return {"degraded": False, "degraded_reason": "", **_read_textlike(path)}


async def _file_read(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    """office.file.read：DOCS_DIR 内文档内容提取 + 溯源（缺文件 404）。"""
    _ = ctx
    filename = str(args.get("filename") or "").strip()
    path = resolve_under_docs(filename, _READ_SUFFIXES)
    if not path.is_file():
        raise BusinessError(
            ErrorCode.NOT_FOUND, f"文件不存在：{path.name}（请先放入文档工作目录）", 404
        )
    if path.stat().st_size > _MAX_FILE_BYTES:
        raise BusinessError(ErrorCode.PARAM_INVALID, "文件过大（超 200KB），请拆分后再解析")
    try:
        result = await asyncio.to_thread(_extract, path)
    except (OSError, ValueError) as exc:
        logger.warning("文件解析失败 %s：%s", path.name, str(exc)[:120])
        raise BusinessError(ErrorCode.PARAM_INVALID, f"文件解析失败：{str(exc)[:120]}") from exc
    return {
        "filename": path.name,
        "format": path.suffix.lower().lstrip("."),
        **result,
        "source": f"local-docs:{path.name}",
    }


async def _docs_rename(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    """office.docs.rename：DOCS_DIR 内批量重命名（演示实现，审批通过后由 decide 路径真正执行）。

    本 handler 仅在审批通过后被调用。目标已存在直接 1001（不覆盖，防丢文件）。
    """
    _ = ctx
    pairs = args.get("pairs")
    if not isinstance(pairs, list) or not pairs:
        raise BusinessError(ErrorCode.PARAM_INVALID, "参数 pairs 必须是非空数组")
    if len(pairs) > _MAX_PAIRS:
        raise BusinessError(ErrorCode.PARAM_INVALID, f"参数 pairs 最多 {_MAX_PAIRS} 对")
    renamed: list[dict[str, str]] = []
    for index, pair in enumerate(pairs):
        if not isinstance(pair, dict):
            raise BusinessError(
                ErrorCode.PARAM_INVALID, f"第 {index + 1} 对必须是键值对（含 src/dst）"
            )
        src = resolve_under_docs(str(pair.get("src") or ""), _RENAME_SUFFIXES)
        dst = resolve_under_docs(str(pair.get("dst") or ""), _RENAME_SUFFIXES)
        if not src.is_file():
            raise BusinessError(
                ErrorCode.NOT_FOUND,
                f"第 {index + 1} 对源文件不存在：{src.name}（重放/误删后重跑会到此，属诚实失败）",
                404,
            )
        if dst.exists():
            raise BusinessError(
                ErrorCode.PARAM_INVALID, f"第 {index + 1} 对目标已存在，拒绝覆盖：{dst.name}"
            )
        src.rename(dst)
        renamed.append({"src": src.name, "dst": dst.name})
    return {
        "renamed": renamed,
        "count": len(renamed),
        "note": "演示实现：DOCS_DIR 内改名，无外部副作用",
    }


def specs() -> tuple[ToolSpec, ...]:
    """文件解析与重命名工具的 ToolSpec（读免审 / 写恒送审）。"""
    return (
        ToolSpec(
            name="office.file.read",
            scope=SCOPE_READ,
            description="提取文档内容：DOCS_DIR 内 docx（含表格文本与图片清单）/xlsx（每表行列数与前 N 行）/csv/txt/md 直读；PDF 引擎缺失如实降级；缺文件 404",
            params={
                "type": "object",
                "properties": {
                    "filename": {
                        "type": "string",
                        "description": "文件名（不含路径，docx/xlsx/csv/txt/md/pdf）",
                        "minLength": 1,
                    },
                },
                "required": ["filename"],
                "additionalProperties": False,
            },
            handler=_file_read,
        ),
        ToolSpec(
            name="office.docs.rename",
            scope=SCOPE_WRITE,
            description="批量重命名 DOCS_DIR 内文件（写动作）：幂等 False（重放因源文件已迁走诚实失败）+ scope=office:write + 恒送审，目标已存在拒绝覆盖",
            params={
                "type": "object",
                "properties": {
                    "pairs": {
                        "type": "array",
                        "description": "改名对数组，元素形如 {src, dst}（均不含路径，最多 20 对）",
                        "items": {
                            "type": "object",
                            "properties": {
                                "src": {"type": "string"},
                                "dst": {"type": "string"},
                            },
                            "required": ["src", "dst"],
                        },
                    },
                },
                "required": ["pairs"],
                "additionalProperties": False,
            },
            handler=_docs_rename,
            idempotent=False,
            requires_approval=True,
            approval_action="office.docs.rename",
        ),
    )


def register_all() -> list[str]:
    """注册文件解析与重命名工具；返回已注册工具名列表。"""
    return [register(spec).name for spec in specs()]
