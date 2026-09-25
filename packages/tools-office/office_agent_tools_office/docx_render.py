"""Word 直出工具（office.docx.render，对齐 PRD §2.1 文档生成「一键导出 Word」）。

职责：
- office.docx.render（office:read 免审）：把标题 + Markdown 子集文本渲染为 .docx 二进制，
  base64 信封直出**不写盘**（与 office.data.export excel 同口径）——对话页对日报/周报/
  纪要等 markdown 产物提供「下载 Word」，前端拿到 base64 还原 Blob 客户端落盘。

链路：__init__.register_all() → registry.register(spec) → executor 直执行（读免审）；
      渲染为纯内存字节，磁盘 IO 零接触，asyncio.to_thread 防大文档阻塞事件循环。
红线：python-docx 缺失则整体不注册（宁缺席，不注册注定调不通的工具）；
      内容全部来自入参原值（只做格式转换不改写文字，数值不可编造）；
      出参带 source + fetched_at 溯源。
对齐：AGENTS.md §3（工具纯函数/降级不 500/溯源标注）；
      智能办公Agent 产品需求文档.md §2.1（智能内容生成与文档处理）。
"""

from __future__ import annotations

import asyncio
import base64
import io
import logging
import re
from datetime import UTC, datetime
from typing import Any

from office_agent_core.contracts import ToolContext, ToolSpec
from office_agent_core.errors import BusinessError, ErrorCode
from office_agent_core.registry import register

try:  # 可选依赖：缺库则本模块不注册任何工具（降级不阻断启动）
    import docx as _docx_module
except ImportError:  # pragma: no cover
    _docx_module = None  # type: ignore[assignment]

logger = logging.getLogger(__name__)

SCOPE_READ = "office:read"

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_MAX_TITLE = 100
_MAX_MARKDOWN = 20000

_HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
_BULLET_RE = re.compile(r"^[-*•]\s+(.*)$")
_ORDERED_RE = re.compile(r"^\d+[.、)]\s*(.*)$")
_BOLD_RE = re.compile(r"\*\*([^*]+)\*\*")


def _now_text() -> str:
    """显式 UTC：溯源时间戳必须可跨时区比对。"""
    return datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S UTC")


def _add_text(paragraph: Any, text: str) -> None:
    """段落内 **加粗** 片段渲染为粗体 run，其余原样（只做格式转换不改写文字）。"""
    last = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > last:
            paragraph.add_run(text[last : match.start()])
        paragraph.add_run(match.group(1)).bold = True
        last = match.end()
    if last < len(text):
        paragraph.add_run(text[last:])


def _table_rows(block: list[str]) -> list[list[str]]:
    """连续 | 行 → 单元格矩阵（分隔行 |---| 剔除，行首尾空管段剔除）。"""
    rows: list[list[str]] = []
    for line in block:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
            continue
        rows.append(cells)
    return rows


def _flush_table(document: Any, block: list[str]) -> None:
    rows = _table_rows(block)
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r_index, row in enumerate(rows):
        for c_index in range(width):
            text = row[c_index] if c_index < len(row) else ""
            cell = table.cell(r_index, c_index)
            paragraph = cell.paragraphs[0]
            if r_index == 0:
                run = paragraph.add_run(text)
                run.bold = True
            else:
                _add_text(paragraph, text)


def render_docx_bytes(title: str, markdown: str) -> bytes:
    """Markdown 子集（标题/列表/表格/加粗/分隔线）→ docx 内存字节，不写盘。"""
    if _docx_module is None:  # pragma: no cover
        raise BusinessError(ErrorCode.PARAM_INVALID, "python-docx 未安装，无法渲染 Word")
    document = _docx_module.Document()
    document.add_heading(title, level=0)
    lines = markdown.splitlines()
    table_block: list[str] = []
    first_heading_skipped = False
    for raw in lines:
        line = raw.strip()
        if line.startswith("|"):
            table_block.append(line)
            continue
        if table_block:
            _flush_table(document, table_block)
            table_block = []
        if not line:
            continue
        heading = _HEADING_RE.match(line)
        if heading:
            # 正文首个一级标题与文档标题重复时跳过（日报常见「# 标题」开头）
            if not first_heading_skipped and len(heading.group(1)) == 1:
                first_heading_skipped = True
                if heading.group(2).strip() == title:
                    continue
            document.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 4))
            continue
        first_heading_skipped = True
        if re.fullmatch(r"-{3,}", line):
            continue
        bullet = _BULLET_RE.match(line)
        if bullet:
            _add_text(document.add_paragraph(style="List Bullet"), bullet.group(1))
            continue
        ordered = _ORDERED_RE.match(line)
        if ordered:
            _add_text(document.add_paragraph(style="List Number"), ordered.group(1))
            continue
        _add_text(document.add_paragraph(), line)
    if table_block:
        _flush_table(document, table_block)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _safe_filename(title: str) -> str:
    """标题作文件名：剔除 Windows/Linux 非法字符，限长防溢出。"""
    cleaned = re.sub(r'[\\/:*?"<>|\r\n]', "_", title).strip(". ")
    return f"{cleaned[:60] or 'document'}.docx"


async def _docx_render(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    """office.docx.render：标题 + markdown → docx base64（读免审不写盘，格式转换零改写）。"""
    _ = ctx
    title = str(args.get("title") or "").strip()
    if not title:
        raise BusinessError(ErrorCode.PARAM_INVALID, "参数 title 不能为空：请传入文档标题")
    if len(title) > _MAX_TITLE:
        raise BusinessError(
            ErrorCode.PARAM_INVALID, f"标题最长 {_MAX_TITLE} 字（当前 {len(title)}）"
        )
    markdown = str(args.get("markdown") or "")
    if not markdown.strip():
        raise BusinessError(ErrorCode.PARAM_INVALID, "参数 markdown 不能为空：请传入文档正文")
    if len(markdown) > _MAX_MARKDOWN:
        raise BusinessError(
            ErrorCode.PARAM_INVALID, f"正文最长 {_MAX_MARKDOWN} 字（当前 {len(markdown)}）"
        )
    raw = await asyncio.to_thread(render_docx_bytes, title, markdown)
    return {
        "filename": _safe_filename(title),
        "content": base64.b64encode(raw).decode("ascii"),
        "encoding": "base64",
        "mime": DOCX_MIME,
        "size_bytes": len(raw),
        "note": "base64 的 .docx 二进制，未写盘；正文为入参 markdown 子集原样渲染（未调大模型）",
        "source": "office.docx.render",
        "fetched_at": _now_text(),
    }


def specs() -> tuple[ToolSpec, ...]:
    """Word 直出工具的 ToolSpec（读免审；依赖缺失返回空）。"""
    if _docx_module is None:  # pragma: no cover
        logger.warning("python-docx 未安装：office.docx.render 不注册（pip install python-docx）")
        return ()
    return (
        ToolSpec(
            name="office.docx.render",
            scope=SCOPE_READ,
            description="Markdown 子集直出 .docx（只读不写盘）：标题+正文（支持标题/列表/表格/"
            "加粗）渲染为 Word 二进制 base64 直出，供对话页「下载 Word」还原 Blob 客户端落盘",
            params={
                "type": "object",
                "properties": {
                    "title": {
                        "type": "string",
                        "description": "文档标题（docx 大标题与文件名）",
                        "minLength": 1,
                        "maxLength": _MAX_TITLE,
                    },
                    "markdown": {
                        "type": "string",
                        "description": "文档正文（Markdown 子集：# 标题 / - 列表 / 1. 编号 / | 表格 | / **加粗**）",
                        "minLength": 1,
                        "maxLength": _MAX_MARKDOWN,
                    },
                },
                "required": ["title", "markdown"],
                "additionalProperties": False,
            },
            handler=_docx_render,
        ),
    )


def register_all() -> list[str]:
    """注册 Word 直出工具（依赖缺失返回空列表）；返回已注册工具名列表。"""
    return [register(spec).name for spec in specs()]
