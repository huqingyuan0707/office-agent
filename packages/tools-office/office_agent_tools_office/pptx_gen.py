"""PPT 生成工具（office.pptx.generate，对齐 PRD §2.12 / §5.3 V1.2「PPT 生成」）。

职责：
- office.pptx.generate（office:write + 恒送审 + idem_key 必填）：把大纲（每页标题 + 要点）
  直出为 .pptx 演示文稿，写入 DOCS_DIR 文档工作目录——「基于文档、周报一键生成 PPT」
  由对话层把 report/minutes 产物整理成大纲后传入，本工具只做模板直出，不调大模型。

链路：__init__.register_all() → registry.register(spec) → executor.call 落审批单，
      复核人批准后由 decide 路径以申请人身份触发 handler 真正写盘；
      磁盘 IO 走 asyncio.to_thread（不阻塞事件循环）。
红线：python-pptx 缺失则整体不注册（宁缺席，不注册注定调不通的工具）；
      文件名经 resolve_under_docs 锁进 DOCS_DIR（basename 化 + 后缀白名单 + 越界拒绝）；
      内容由入参原值直出（数值不可编造）；写动作恒送审，幂等键必带。
对齐：AGENTS.md §3（写动作恒送审/数据不出域在本地盘的对应实现）；
      智能办公Agent 产品需求文档.md §2.12（多模态：一键生成 PPT 大纲与每页文案）。
"""

from __future__ import annotations

import asyncio
import logging
from datetime import UTC, datetime
from typing import Any

from office_agent_core.contracts import ToolContext, ToolSpec
from office_agent_core.errors import BusinessError, ErrorCode
from office_agent_core.registry import register

from .paths import resolve_under_docs

try:  # 可选依赖：缺库则本模块不注册任何工具（降级不阻断启动）
    from pptx import Presentation
except ImportError:  # pragma: no cover
    Presentation = None  # type: ignore[assignment]

logger = logging.getLogger(__name__)

SCOPE_WRITE = "office:write"

_MAX_SLIDES = 30
_MAX_BULLETS = 20
_TITLE_MAX = 60
_BULLET_MAX = 200


def _now_text() -> str:
    """显式 UTC：溯源时间戳必须可跨时区比对。"""
    return datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S UTC")


def _slides_or_raise(raw: Any) -> list[dict[str, Any]]:
    """大纲口径：1-30 页，每页标题 1-60 字 + 要点 1-20 条（每条 1-200 字，不臆造内容）。"""
    if not isinstance(raw, list) or not raw:
        raise BusinessError(
            ErrorCode.PARAM_INVALID,
            '参数 slides 必须为非空数组，元素形如 {"title": "页标题", "bullets": ["要点1", "要点2"]}',
        )
    if len(raw) > _MAX_SLIDES:
        raise BusinessError(
            ErrorCode.PARAM_INVALID, f"PPT 最多 {_MAX_SLIDES} 页（当前 {len(raw)} 页）"
        )
    slides: list[dict[str, Any]] = []
    for index, item in enumerate(raw, 1):
        if not isinstance(item, dict):
            raise BusinessError(
                ErrorCode.PARAM_INVALID, f"第 {index} 页必须是对象（title + bullets）"
            )
        title = str(item.get("title") or "").strip()
        if not title or len(title) > _TITLE_MAX:
            raise BusinessError(
                ErrorCode.PARAM_INVALID, f"第 {index} 页标题必须为 1-{_TITLE_MAX} 字的非空文本"
            )
        bullets_raw = item.get("bullets")
        if not isinstance(bullets_raw, list) or not bullets_raw:
            raise BusinessError(
                ErrorCode.PARAM_INVALID, f"第 {index} 页（{title}）要点必须为非空数组"
            )
        if len(bullets_raw) > _MAX_BULLETS:
            raise BusinessError(
                ErrorCode.PARAM_INVALID,
                f"第 {index} 页（{title}）要点最多 {_MAX_BULLETS} 条（当前 {len(bullets_raw)} 条）",
            )
        bullets = [str(bullet).strip() for bullet in bullets_raw]
        if any(not bullet or len(bullet) > _BULLET_MAX for bullet in bullets):
            raise BusinessError(
                ErrorCode.PARAM_INVALID,
                f"第 {index} 页（{title}）每条要点必须为 1-{_BULLET_MAX} 字的非空文本",
            )
        slides.append({"title": title, "bullets": bullets})
    return slides


async def _pptx_generate(ctx: ToolContext, args: dict[str, Any]) -> dict[str, Any]:
    """office.pptx.generate：大纲直出 .pptx（封面标题 + 每页要点；审批通过后才真正写盘）。"""
    _ = ctx
    title = str(args.get("title") or "").strip()
    if not title:
        raise BusinessError(ErrorCode.PARAM_INVALID, "参数 title 不能为空：请传入演示文稿标题")
    slides = _slides_or_raise(args.get("slides"))
    path = resolve_under_docs(str(args.get("filename") or ""), (".pptx",))

    def _write() -> int:
        deck = Presentation()
        cover = deck.slides.add_slide(deck.slide_layouts[0])
        cover.shapes.title.text = title
        for slide_spec in slides:
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = slide_spec["title"]
            body = slide.placeholders[1].text_frame
            for index, bullet in enumerate(slide_spec["bullets"]):
                paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
                paragraph.text = bullet
        deck.save(str(path))
        return len(slides)

    count = await asyncio.to_thread(_write)
    return {
        "file": path.name,
        "title": title,
        "slide_count": count,
        "published": True,
        "published_at": _now_text(),
        "note": "演示文稿已写入文档工作目录；内容为入参大纲原值直出（未调大模型）",
    }


def specs() -> tuple[ToolSpec, ...]:
    """PPT 生成工具的 ToolSpec（写口径：恒送审 + 幂等键必带；依赖缺失返回空）。"""
    if Presentation is None:  # pragma: no cover
        logger.warning("python-pptx 未安装：office.pptx.generate 不注册（pip install python-pptx）")
        return ()
    return (
        ToolSpec(
            name="office.pptx.generate",
            scope=SCOPE_WRITE,
            description="生成 .pptx 演示文稿（写动作）：大纲每页标题+要点直出（不调大模型），"
            "恒送审 + idem_key 必填，审批通过后才真正写盘到文档工作目录；"
            "基于周报/纪要生成 PPT 时由对话层先整理大纲再传入",
            params={
                "type": "object",
                "properties": {
                    "title": {
                        "type": "string",
                        "description": "演示文稿标题（同时作为封面标题）",
                        "minLength": 1,
                        "maxLength": 100,
                    },
                    "filename": {
                        "type": "string",
                        "description": "目标文件名（以 .pptx 结尾，不含路径）",
                        "minLength": 6,
                        "maxLength": 120,
                    },
                    "slides": {
                        "type": "array",
                        "description": '大纲数组（1-30 页），元素形如 {"title": "页标题", "bullets": ["要点"]}',
                        "items": {
                            "type": "object",
                            "properties": {
                                "title": {"type": "string"},
                                "bullets": {"type": "array", "items": {"type": "string"}},
                            },
                            "required": ["title", "bullets"],
                        },
                    },
                    "idem_key": {
                        "type": "string",
                        "description": "幂等键（8-64 字符）",
                        "minLength": 8,
                        "maxLength": 64,
                    },
                },
                "required": ["title", "filename", "slides", "idem_key"],
                "additionalProperties": False,
            },
            handler=_pptx_generate,
            idempotent=True,
            requires_approval=True,
            approval_action="office.pptx.generate",
        ),
    )


def register_all() -> list[str]:
    """注册 PPT 生成工具（依赖缺失返回空列表）；返回已注册工具名列表。"""
    return [register(spec).name for spec in specs()]
