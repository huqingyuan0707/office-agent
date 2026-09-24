"""V1.2 批次 A HTTP 级冒烟（PPT 生成 / 合规检测 / 预算查询，可重复执行）：

① 工具清单含三新工具且口径正确（pptx=office:write 恒送审；scan/query=office:read）；
② office.compliance.scan：邮件场景命中手机号+绝对化用语+凭据脱敏；干净文本 passed；
③ office.budget.query：全量台账 + 项目过滤 + 超支/紧张状态 + 汇总口径；
④ office.pptx.generate：invoke 落审批单（pending_approval）→ reviewer 批准 →
   以申请人身份写盘 DOCS_DIR，文件真实存在且 .pptx 可读；
⑤ 穿越文件名被 1001 拒绝（数据不出域在本地盘的对应实现）。

运行前提：`.venv\\Scripts\\python.exe -m office_agent_server`（8200）已在跑。
"""

import sys
import uuid
from pathlib import Path

import httpx
from pptx import Presentation

from office_agent_core.settings import settings

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
OFFICE = "http://127.0.0.1:8200/api/v1"

failures: list[str] = []


def check(step: str, cond: bool, detail: str = "") -> None:
    print(
        ("PASS | " if cond else "FAIL | ") + step + (f"（{detail}）" if detail and not cond else "")
    )
    if not cond:
        failures.append(step)


def login(c: httpx.Client, u: str, p: str) -> dict:
    r = c.post("/auth/login", json={"username": u, "password": p}).json()
    return {"Authorization": f"Bearer {r['data']['token']}"}


def payload(r: dict) -> dict:
    """invoke 信封出参在 data.result（data 外层是 tool/status/args 等元信息）。"""
    return (r.get("data") or {}).get("result") or {}


with httpx.Client(base_url=OFFICE, trust_env=False, timeout=30) as c:
    admin = login(c, "admin", "admin123")
    reviewer = login(c, "reviewer", "reviewer123")

    # ① 工具清单
    tools = {t["name"]: t for t in c.get("/agent/tools", headers=admin).json()["data"]["items"]}
    pptx_spec, scan_spec, budget_spec = (
        tools.get("office.pptx.generate"),
        tools.get("office.compliance.scan"),
        tools.get("office.budget.query"),
    )
    check(
        "① 三新工具已注册且口径正确",
        pptx_spec is not None
        and pptx_spec["requires_approval"] is True
        and pptx_spec["scope"] == "office:write"
        and scan_spec is not None
        and scan_spec["requires_approval"] is False
        and scan_spec["scope"] == "office:read"
        and budget_spec is not None
        and budget_spec["scope"] == "office:read",
        str(
            {
                k: v and (v["scope"], v["requires_approval"])
                for k, v in tools.items()
                if k.startswith("office.pptx")
                or k.startswith("office.compliance")
                or k.startswith("office.budget")
            }
        ),
    )

    # ② 合规扫描
    r = c.post(
        "/agent/tools/office.compliance.scan/invoke",
        headers=admin,
        json={
            "args": {
                "text": "联系张三 13812345678，本品全网最低价。password=SuperSecret123",
                "scene": "email",
            }
        },
    ).json()
    d = payload(r)
    check(
        "②a 合规扫描三类命中（隐私/用语/凭据脱敏）",
        r.get("code") == 0
        and d.get("passed") is False
        and d.get("counts", {}).get("privacy") == 1
        and d.get("counts", {}).get("words") >= 1
        and d.get("counts", {}).get("leak") == 1
        and all("SuperSecret123" not in h["excerpt"] for h in d.get("hits", [])),
        str(d)[:300],
    )
    r = c.post(
        "/agent/tools/office.compliance.scan/invoke",
        headers=admin,
        json={"args": {"text": "本周例会按期推进，结论已同步。"}},
    ).json()
    check(
        "②b 干净文本 passed=True",
        r.get("code") == 0 and payload(r).get("passed") is True,
        str(r)[:200],
    )

    # ③ 预算查询
    r = c.post("/agent/tools/office.budget.query/invoke", headers=admin, json={"args": {}}).json()
    d = payload(r)
    status_map = {row["project"]: row["status"] for row in d.get("rows", [])}
    check(
        "③a 预算台账全量（溯源+汇总+状态标签）",
        r.get("code") == 0
        and d.get("source") == "builtin-demo"
        and d.get("count") == 5
        and status_map.get("数据看板") == "超支"
        and status_map.get("客服知识库") == "紧张"
        and d.get("summary", {}).get("total_remaining")
        == d.get("summary", {}).get("total_budget", 0) - d.get("summary", {}).get("total_used", 0),
        str(d)[:300],
    )
    r = c.post(
        "/agent/tools/office.budget.query/invoke", headers=admin, json={"args": {"project": "改版"}}
    ).json()
    d = payload(r)
    check(
        "③b 项目过滤命中官网改版（remaining=74000）",
        r.get("code") == 0
        and d.get("count") == 1
        and d["rows"][0]["project"] == "官网改版"
        and d["rows"][0]["remaining"] == 74000,
        str(d)[:200],
    )

    # ④ PPT 生成：恒送审 → 批准 → 写盘
    filename = f"smoke-v12-{uuid.uuid4().hex[:8]}.pptx"
    r = c.post(
        "/agent/tools/office.pptx.generate/invoke",
        headers=admin,
        json={
            "args": {
                "title": "周报复盘",
                "filename": filename,
                "slides": [
                    {"title": "本周进展", "bullets": ["日报助手上线", "审批闭环跑通"]},
                    {"title": "下周计划", "bullets": ["V1.2 批次 B 前端两页"]},
                ],
                "idem_key": f"smoke-pptx-{uuid.uuid4().hex[:12]}",
            }
        },
    ).json()
    d = r.get("data") or {}
    check(
        "④a pptx.generate invoke 只落审批单（pending_approval）",
        r.get("code") == 0
        and d.get("status") == "pending_approval"
        and d.get("approval_required") is True,
        str(d)[:300],
    )
    r2 = c.post(
        f"/approvals/{d['approval_id']}/approve", headers=reviewer, json={"reason": "V1.2 冒烟"}
    ).json()
    ex = (r2.get("data") or {}).get("execution_result") or {}
    path = Path(settings.DOCS_DIR) / filename
    readable = False
    if path.exists():
        deck = Presentation(str(path))
        texts = "\n".join(
            shape.text_frame.text
            for slide in deck.slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
        readable = "周报复盘" in texts and "日报助手上线" in texts
    check(
        "④b 复核员批准后真实写盘且内容可读",
        r2.get("code") == 0 and ex.get("status") == "ok" and path.exists() and readable,
        str(ex)[:300],
    )

    # ⑤ 穿越拒绝（写动作 invoke 只落单；路径校验在批准后执行期触发 → 驳回流）
    r = c.post(
        "/agent/tools/office.pptx.generate/invoke",
        headers=admin,
        json={
            "args": {
                "title": "穿越",
                "filename": "../evil.pptx",
                "slides": [{"title": "页", "bullets": ["点"]}],
                "idem_key": f"smoke-pptx-{uuid.uuid4().hex[:12]}",
            }
        },
    ).json()
    d = r.get("data") or {}
    check(
        "⑤a 穿越请求先正常落审批单",
        r.get("code") == 0 and d.get("status") == "pending_approval",
        str(d)[:200],
    )
    r2 = c.post(
        f"/approvals/{d['approval_id']}/approve", headers=reviewer, json={"reason": "穿越验证"}
    ).json()
    ex = (r2.get("data") or {}).get("execution_result") or {}
    evil_outside = (Path(settings.DOCS_DIR).resolve().parent / "evil.pptx").exists()
    check(
        "⑤b 批准后执行期 1001 拒绝且未越界写盘",
        ex.get("status") == "failed" and ex.get("code") == 1001 and not evil_outside,
        str(ex)[:200],
    )

print()
print("RESULT:", "ALL PASS" if not failures else f"{len(failures)} FAILED: {failures}")
sys.exit(1 if failures else 0)
