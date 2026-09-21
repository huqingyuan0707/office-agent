"""办公文档工具包（可选能力：python-docx / openpyxl / python-pptx 缺哪个就少注册对应工具）。

职责：注册四个文档处理工具——
① office.docx.read（office:read）：解析 .docx → 段落与表格结构化内容，带来源与计数；
② office.xlsx.read（office:read）：读取 .xlsx 工作表 → 行数据（max_rows 截断保护）；
③ office.docx.write（office:write + needs_approval=True）：生成 .docx（标题 + 段落），过审批中心才执行；
④ office.pptx.write（office:write + needs_approval=True）：生成 .pptx（标题 + 每页要点），同走审批。

红线：
- 写动作无直接生效通道：invoke 只落审批单，复核人批准后由 decide 路径真正写盘（与 office.memo.submit 同一套双人闭环）；
- 内容由入参原值直出（数值不编造）；读响应必带 source 与 extracted_at 溯源；
- 路径安全：读写都锁在 DOCS_DIR 工作目录内（basename 化 + resolve 前缀校验 + 后缀白名单），绝不越界。

链路：main.lifespan → register_all() → registry；executor 经 ToolSpec.handler(args) 调用；
磁盘 IO 走 asyncio.to_thread（不阻塞事件循环）。
对齐：docs/office-agent仓库骨架与内核提取方案.md §4（tools 层扩展：办公文档域）。
"""

from __future__ import annotations

import asyncio
import logging
import os
from datetime import datetime, timezone
from pathlib import Path

from office_agent import registry
from office_agent.config import settings
from office_agent.contracts import SCOPE_READ, SCOPE_WRITE, ToolError, ToolSpec

try:  # 可选依赖：缺哪个库就少注册对应工具（宁缺席，不注册注定调不通的工具）
    import docx  # python-docx
except ImportError:  # pragma: no cover
    docx = None  # type: ignore[assignment]

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None  # type: ignore[assignment]

try:
    from pptx import Presentation  # python-pptx
except ImportError:  # pragma: no cover
    Presentation = None  # type: ignore[assignment]

logger = logging.getLogger(__name__)

_SUFFIXES = {"docx": ".docx", "xlsx": ".xlsx", "pptx": ".pptx"}


def _now_text() -> str:
    """显式 UTC 溯源时间戳（跨时区可比对）。"""
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def _resolve_under_docs(filename: str, kind: str) -> Path:
    """把入参文件名锁进 DOCS_DIR：basename 化防穿越 + 后缀白名单，返回可用的绝对路径。"""
    suffix = _SUFFIXES[kind]
    root = Path(settings.DOCS_DIR).resolve()
    root.mkdir(parents=True, exist_ok=True)
    raw = str(filename or "").strip()
    name = os.path.basename(raw)
    if not name:
        raise ToolError(
            400, f"参数 filename 不能为空：请传入 {suffix} 文件名（仅文件名，不含路径）"
        )
    if name != raw or ".." in raw:  # 含路径分隔符或相对段：显式拒绝（比静默剥离更可诊断）
        raise ToolError(400, f"文件名只能是不含路径的名称（拒绝：{raw}）")
    if not name.lower().endswith(suffix):
        raise ToolError(400, f"文件名必须以 {suffix} 结尾（当前：{name}）")
    target = (root / name).resolve()
    if root not in target.parents:
        raise ToolError(400, "文件名非法：不允许路径分隔符或越界访问（只能写入文档工作目录）")
    return target


def _source_tag(path: Path) -> str:
    """读侧溯源标注：local-docs:<文件名>（目录内相对口径）。"""
    return f"local-docs:{path.name}"


async def _docx_read(args: dict) -> dict:
    """office.docx.read：解析段落与表格，内容原值直出并标注来源。"""
    if docx is None:
        raise ToolError(503, "python-docx 未安装：pip install python-docx 后重启服务")
    path = _resolve_under_docs(str(args.get("path") or ""), "docx")
    if not path.exists():
        raise ToolError(404, f"文件不存在：{path.name}（文档工作目录 {settings.DOCS_DIR}）")

    def _parse() -> dict:
        document = docx.Document(str(path))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        tables: list[list[list[str]]] = []
        for table in document.tables:
            tables.append([[cell.text.strip() for cell in row.cells] for row in table.rows])
        return {"paragraphs": paragraphs, "tables": tables}

    parsed = await asyncio.to_thread(_parse)
    return {
        "file": path.name,
        "paragraphs": parsed["paragraphs"],
        "tables": parsed["tables"],
        "counts": {
            "paragraph": len(parsed["paragraphs"]),
            "table": len(parsed["tables"]),
        },
        "source": _source_tag(path),
        "extracted_at": _now_text(),
    }


async def _xlsx_read(args: dict) -> dict:
    """office.xlsx.read：读工作表行数据（表头 + 数据行，max_rows 截断保护）。"""
    if openpyxl is None:
        raise ToolError(503, "openpyxl 未安装：pip install openpyxl 后重启服务")
    path = _resolve_under_docs(str(args.get("path") or ""), "xlsx")
    if not path.exists():
        raise ToolError(404, f"文件不存在：{path.name}（文档工作目录 {settings.DOCS_DIR}）")
    max_rows = args.get("max_rows") or 50
    if not isinstance(max_rows, int) or isinstance(max_rows, bool) or not 1 <= max_rows <= 500:
        raise ToolError(400, "参数 max_rows 必须是 1-500 的整数（防大表拖垮响应）")

    def _parse() -> tuple[str, list[str], list[list], bool]:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            sheet_name = str(args.get("sheet") or workbook.sheetnames[0])
            if sheet_name not in workbook.sheetnames:
                raise ToolError(
                    400, f"工作表「{sheet_name}」不存在，可选：{' / '.join(workbook.sheetnames)}"
                )
            sheet = workbook[sheet_name]
            rows_raw = [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
        truncated = len(rows_raw) > max_rows + 1  # 首行按表头计
        rows_raw = rows_raw[: max_rows + 1]
        if not rows_raw:
            return sheet_name, [], [], False
        header = [str(cell) if cell is not None else "" for cell in rows_raw[0]]
        data = [list(row) for row in rows_raw[1:]]
        return sheet_name, header, data, truncated

    sheet_name, header, rows, truncated = await asyncio.to_thread(_parse)
    return {
        "file": path.name,
        "sheet": sheet_name,
        "header": header,
        "rows": rows,
        "counts": {"row": len(rows), "column": len(header)},
        "truncated": truncated,
        "source": _source_tag(path),
        "extracted_at": _now_text(),
    }


async def _docx_write(args: dict) -> dict:
    """office.docx.write：生成 .docx（标题 + 段落，入参原值直出；审批通过后才被 decide 触发）。"""
    if docx is None:
        raise ToolError(503, "python-docx 未安装：pip install python-docx 后重启服务")
    path = _resolve_under_docs(str(args.get("filename") or ""), "docx")
    title = str(args.get("title") or "").strip()
    paragraphs = args.get("paragraphs")
    if not title:
        raise ToolError(400, "参数 title 不能为空：请传入文档标题")
    if (
        not isinstance(paragraphs, list)
        or not paragraphs
        or not all(isinstance(p, str) and p.strip() for p in paragraphs)
    ):
        raise ToolError(400, '参数 paragraphs 必须为非空字符串数组，形如 ["第一段", "第二段"]')

    def _write() -> int:
        document = docx.Document()
        document.add_heading(title, level=0)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph.strip())
        document.save(str(path))
        return len(paragraphs)

    count = await asyncio.to_thread(_write)
    return {
        "file": path.name,
        "title": title,
        "paragraph_count": count,
        "published": True,
        "published_at": _now_text(),
        "note": "演示实现：内容入参原值直出，仅写入文档工作目录；真实接入时在此对接存储/通知渠道",
    }


async def _pptx_write(args: dict) -> dict:
    """office.pptx.write：生成 .pptx（封面标题 + 每页要点列表；审批通过后才真正写盘）。"""
    if Presentation is None:
        raise ToolError(503, "python-pptx 未安装：pip install python-pptx 后重启服务")
    path = _resolve_under_docs(str(args.get("filename") or ""), "pptx")
    title = str(args.get("title") or "").strip()
    slides = args.get("slides")
    if not title:
        raise ToolError(400, "参数 title 不能为空：请传入演示文稿标题")
    if (
        not isinstance(slides, list)
        or not slides
        or not all(
            isinstance(s, dict)
            and str(s.get("title") or "").strip()
            and isinstance(s.get("bullets"), list)
            and all(isinstance(b, str) and b.strip() for b in s["bullets"])
            for s in slides
        )
    ):
        raise ToolError(
            400,
            '参数 slides 必须为非空数组，元素形如 {"title": "页标题", "bullets": ["要点1", "要点2"]}',
        )

    def _write() -> int:
        deck = Presentation()
        cover = deck.slides.add_slide(deck.slide_layouts[0])
        cover.shapes.title.text = title
        for slide_spec in slides:
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = str(slide_spec["title"]).strip()
            body = slide.placeholders[1].text_frame
            for index, bullet in enumerate(slide_spec["bullets"]):
                paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
                paragraph.text = bullet.strip()
        deck.save(str(path))
        return len(slides)

    count = await asyncio.to_thread(_write)
    return {
        "file": path.name,
        "title": title,
        "slide_count": count,
        "published": True,
        "published_at": _now_text(),
        "note": "演示实现：内容入参原值直出，仅写入文档工作目录；真实接入时在此对接模板/品牌规范",
    }


def register_all() -> None:
    """按依赖可用性注册文档工具（幂等；重复调用覆盖同规格项，重名抛 ToolError 直接暴露问题）。"""
    specs: tuple[ToolSpec, ...] = ()
    if docx is not None:
        specs += (
            ToolSpec(
                name="office.docx.read",
                description="解析 .docx 文档：返回段落与表格结构化内容（带来源与计数，仅限文档工作目录内）",
                scope=SCOPE_READ,
                needs_approval=False,
                schema={
                    "type": "object",
                    "properties": {
                        "path": {"type": "string", "description": "docx 文件名（不含路径）"}
                    },
                    "required": ["path"],
                },
                handler=_docx_read,
            ),
            ToolSpec(
                name="office.docx.write",
                description="生成 .docx 文档（标题 + 段落，入参原值直出）：需审批，审批通过后才真正写盘",
                scope=SCOPE_WRITE,
                needs_approval=True,
                schema={
                    "type": "object",
                    "properties": {
                        "filename": {
                            "type": "string",
                            "description": "目标文件名（以 .docx 结尾，不含路径）",
                        },
                        "title": {"type": "string", "description": "文档标题"},
                        "paragraphs": {
                            "type": "array",
                            "items": {"type": "string"},
                            "description": "段落内容数组",
                        },
                    },
                    "required": ["filename", "title", "paragraphs"],
                },
                handler=_docx_write,
            ),
        )
    if openpyxl is not None:
        specs += (
            ToolSpec(
                name="office.xlsx.read",
                description="读取 .xlsx 工作表：返回表头与数据行（max_rows 截断保护，带来源标注）",
                scope=SCOPE_READ,
                needs_approval=False,
                schema={
                    "type": "object",
                    "properties": {
                        "path": {"type": "string", "description": "xlsx 文件名（不含路径）"},
                        "sheet": {"type": "string", "description": "工作表名（缺省取第一个）"},
                        "max_rows": {
                            "type": "integer",
                            "description": "数据行上限（1-500，默认 50）",
                        },
                    },
                    "required": ["path"],
                },
                handler=_xlsx_read,
            ),
        )
    if Presentation is not None:
        specs += (
            ToolSpec(
                name="office.pptx.write",
                description="生成 .pptx 演示文稿（封面 + 每页要点，入参原值直出）：需审批，审批通过后才真正写盘",
                scope=SCOPE_WRITE,
                needs_approval=True,
                schema={
                    "type": "object",
                    "properties": {
                        "filename": {
                            "type": "string",
                            "description": "目标文件名（以 .pptx 结尾，不含路径）",
                        },
                        "title": {"type": "string", "description": "演示文稿标题"},
                        "slides": {
                            "type": "array",
                            "description": '页数组，元素形如 {"title": "页标题", "bullets": ["要点"]}',
                            "items": {
                                "type": "object",
                                "properties": {
                                    "title": {"type": "string"},
                                    "bullets": {"type": "array", "items": {"type": "string"}},
                                },
                                "required": ["title", "bullets"],
                            },
                        },
                    },
                    "required": ["filename", "title", "slides"],
                },
                handler=_pptx_write,
            ),
        )
    for spec in specs:
        registry.register(spec)
    registered = [spec.name for spec in specs]
    logger.info(
        "文档工具注册完成：%s", "、".join(registered) if registered else "（依赖缺失，跳过）"
    )
